"""
MMCC Site Document Extractor  —  NO API KEY • NO ENDPOINT • NO EXTRA SOFTWARE
----------------------------------------------------------------------------
Tab 1 – Material Inventory : packing slips / delivery tickets / invoices
Tab 2 – Time & Material    : trade T&M tickets / timesheets
Tab 3 – Summary            : dashboard + Excel summary sheets

Reads PHOTOS (HEIC/HEIF, JPG, PNG, WEBP) and PDFs.
 • Digital PDFs           -> text read directly from the PDF
 • Photos / scanned PDFs  -> built-in offline text reader (RapidOCR)

Install & run (pip only):
    pip install -r requirements.txt
    streamlit run mmcc_extractor_no_api.py

requirements.txt
    streamlit>=1.50
    pandas>=2.0
    numpy
    openpyxl>=3.1
    pillow>=10.0
    pillow-heif>=0.16
    pymupdf>=1.24
    rapidocr-onnxruntime>=1.3

Fields are filled by matching labels ("PO #", "Qty", "Approved By"...) and
table columns. Always review the table against the photo before downloading.
"""

import io
import re
from datetime import date, datetime

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageOps

try:
    from pillow_heif import register_heif_opener
    register_heif_opener()
    HEIC_OK = True
except ImportError:
    HEIC_OK = False

try:
    import fitz  # PyMuPDF
    PDF_OK = True
except ImportError:
    PDF_OK = False

FILE_TYPES = ["heic", "heif", "jpg", "jpeg", "png", "webp", "pdf"]
MAX_PDF_PAGES = 10

# ============================================================================
# COLUMN TEMPLATES
# ============================================================================
INV_COLUMNS = [
    "Doc Date", "Date Received", "Vendor", "Doc Type", "Doc #", "PO #",
    "Ordered By", "Product #", "Description", "Category", "Qty", "U/M",
    "Unit Price", "Extension", "Notes",
]
TM_COLUMNS = [
    "Work Date", "Ticket Date", "Trade Company", "T&M Ticket #", "Scope of Work",
    "Reference (RFI/SI/CD)", "Location", "Line Type", "Worker Name", "Classification",
    "Reg Hrs", "OT Hrs", "Total Hrs", "Material / Equipment", "Qty", "U/M",
    "Unit Rate", "Amount", "Trade Signed By", "Approved By (Bird)", "Signed Date", "Notes",
]
DOC_TYPES = ["Packing Slip", "Delivery Ticket", "Invoice", "Bill of Lading",
             "Quote", "Receipt", "Other"]
LINE_TYPES = ["Labour", "Material", "Equipment"]
CATEGORY_KEYWORDS = {
    "Concrete & Rebar": ["rebar", "concrete", "dowel", "grout", "chair", "mesh", "form"],
    "Lumber & Framing": ["lvl", "lumber", "stud", "plywood", "osb", "2x", "joist", "psl"],
    "Steel & Fasteners": ["hss", "bolt", "anchor", "screw", "nail", "washer", "nut", "angle", "plate"],
    "Building Envelope": ["membrane", "avb", "tyvek", "sheathing", "insulation", "flashing",
                          "waterproof", "drainage", "termination bar", "sealant", "caulk"],
    "Drywall & Finishes": ["drywall", "gypsum", "mud", "tape", "corner bead", "paint", "primer"],
    "Mechanical": ["pipe", "fitting", "valve", "duct", "fan coil", "pex", "pvc", "abs"],
    "Electrical": ["wire", "conduit", "breaker", "cable", "emt", "fixture"],
    "Safety & PPE": ["glove", "vest", "hard hat", "glasses", "harness", "barricade", "cone"],
    "Consumables & Tools": ["blade", "bit", "tarp", "poly", "rag", "fuel", "battery"],
}
CATEGORIES = list(CATEGORY_KEYWORDS.keys()) + ["Other"]

UOMS = {"EA", "EACH", "PC", "PCS", "PCE", "BOX", "BX", "LF", "FT", "M", "LM", "ROLL", "RL",
        "BDL", "BUNDLE", "KG", "LB", "LBS", "BAG", "BG", "CTN", "PKG", "PK", "SHT", "SHEET",
        "TUBE", "TB", "PAIL", "GAL", "L", "SET", "LOT", "CS", "CASE", "PR", "SF", "SQFT",
        "M2", "M3", "CY", "YD", "TON", "T", "HR", "HRS", "DAY", "DAYS", "WK", "MO"}
CLASSIFICATIONS = r"(general\s*)?foreman|lead\s*hand|journeyman|journeyperson|j/?m|apprentice|" \
                  r"\d(st|nd|rd|th)\s*y(ea)?r|labou?rer|helper|operator|carpenter|electrician|" \
                  r"plumber|pipefitter|ironworker|finisher|cement\s*mason|welder|supervisor|" \
                  r"superintendent|driver|teamster|glazier|insulator|roofer|painter|taper"

# ============================================================================
# PAGE SETUP
# ============================================================================
st.set_page_config(page_title="MMCC Site Document Extractor", page_icon="🏗️", layout="wide")
st.markdown(
    """
    <style>
    .block-container {padding-top: 1.5rem;}
    .title-bar {background:#003B71;color:white;padding:14px 20px;border-radius:8px;margin-bottom:12px;}
    .title-bar h2 {margin:0;color:white;}
    .title-bar p {margin:2px 0 0 0;color:#D0DCEA;font-size:0.9rem;}
    div[data-testid="stMetricValue"] {color:#003B71;}
    </style>
    <div class="title-bar">
      <h2>🏗️ MMCC Site Document Extractor</h2>
      <p>Material Inventory • Time &amp; Material Tracking • Summary — HEIC / JPG / PNG / PDF — no API key</p>
    </div>
    """,
    unsafe_allow_html=True,
)

for key, cols in (("inv", INV_COLUMNS), ("tm", TM_COLUMNS)):
    st.session_state.setdefault(f"{key}_base", pd.DataFrame(columns=cols + ["Source File"]))
    st.session_state.setdefault(f"{key}_view", st.session_state[f"{key}_base"])
    st.session_state.setdefault(f"{key}_processed", set())
    st.session_state.setdefault(f"{key}_ver", 0)
    st.session_state.setdefault(f"{key}_images", {})
    st.session_state.setdefault(f"{key}_text", {})


# ============================================================================
# 1. READ FILES -> ROWS OF TEXT CELLS
# ============================================================================
@st.cache_resource(show_spinner="Loading text reader (first run only)…")
def get_ocr():
    from rapidocr_onnxruntime import RapidOCR
    return RapidOCR()


def open_image(file, max_side=2400) -> Image.Image:
    file.seek(0)
    img = ImageOps.exif_transpose(Image.open(file)).convert("RGB")  # fix sideways phone photos
    img.thumbnail((max_side, max_side))
    return img


def boxes_to_rows(boxes):
    """Group text boxes into lines, then split each line into cells at big gaps."""
    if not boxes:
        return []
    med = float(np.median([b["y1"] - b["y0"] for b in boxes])) or 10
    for b in boxes:
        b["yc"] = (b["y0"] + b["y1"]) / 2
    lines = []
    for b in sorted(boxes, key=lambda b: b["yc"]):
        if lines and abs(lines[-1]["yc"] - b["yc"]) < med * 0.6:
            lines[-1]["items"].append(b)
            lines[-1]["yc"] = float(np.mean([i["yc"] for i in lines[-1]["items"]]))
        else:
            lines.append({"yc": b["yc"], "items": [b]})
    rows = []
    for ln in lines:
        items = sorted(ln["items"], key=lambda b: b["x0"])
        cur = {"t": items[0]["text"], "x0": items[0]["x0"], "x1": items[0]["x1"]}
        cells = []
        for b in items[1:]:
            if b["x0"] - cur["x1"] > med * 1.2:
                cells.append(cur)
                cur = {"t": b["text"], "x0": b["x0"], "x1": b["x1"]}
            else:
                cur["t"] += " " + b["text"]
                cur["x1"] = max(cur["x1"], b["x1"])
        cells.append(cur)
        rows.append(cells)
    return rows


def ocr_image(img: Image.Image):
    result, _ = get_ocr()(np.array(img))
    boxes = []
    for pts, text, _score in result or []:
        xs, ys = [p[0] for p in pts], [p[1] for p in pts]
        boxes.append({"x0": min(xs), "x1": max(xs), "y0": min(ys), "y1": max(ys), "text": text.strip()})
    return boxes_to_rows(boxes)


def to_jpg(img, max_side=1600):
    img = img.copy()
    img.thumbnail((max_side, max_side))
    b = io.BytesIO()
    img.save(b, "JPEG", quality=85)
    return b.getvalue()


def read_file(file):
    """Return (rows, preview_jpgs). rows = list of [ {t, x0, x1}, ... ]."""
    file.seek(0)
    rows, previews = [], []
    if file.name.lower().endswith(".pdf"):
        if not PDF_OK:
            raise RuntimeError("PDF support not installed – run: pip install pymupdf")
        doc = fitz.open(stream=file.read(), filetype="pdf")
        for i, page in enumerate(doc):
            if i >= MAX_PDF_PAGES:
                break
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            img = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
            words = page.get_text("words")
            if len(words) > 5:  # digital PDF – use its text layer
                rows += boxes_to_rows([{"x0": w[0] * 2, "y0": w[1] * 2, "x1": w[2] * 2,
                                        "y1": w[3] * 2, "text": w[4]} for w in words])
            else:               # scanned PDF – read the picture
                rows += ocr_image(img)
            previews.append(to_jpg(img))
    else:
        img = open_image(file)
        rows = ocr_image(img)
        previews.append(to_jpg(img))
    return rows, previews


def thumb(file):
    file.seek(0)
    if file.name.lower().endswith(".pdf"):
        if not PDF_OK:
            return None
        doc = fitz.open(stream=file.read(), filetype="pdf")
        img = Image.open(io.BytesIO(doc[0].get_pixmap().tobytes("png")))
    else:
        img = open_image(file, 500)
    file.seek(0)
    return img


# ============================================================================
# 2. TEXT HELPERS
# ============================================================================
def rt(row):
    return " ".join(c["t"] for c in row)


def clean(s):
    s = re.sub(r"_{2,}", " ", str(s))
    return re.sub(r"\s+", " ", s).strip(" :#.-|\t")


def find_value(rows, labels, start=0, end=None):
    """Find the first label (priority order). Value = text after label,
    else next cell, else the cell below."""
    for lab in labels:
        pat = re.compile(r"(?:%s)\s*[:#.\-]*\s*(?P<val>.*)$" % lab, re.I)
        for i in range(start, len(rows) if end is None else min(end, len(rows))):
            row = rows[i]
            for j, c in enumerate(row):
                m = pat.search(c["t"])
                if not m:
                    continue
                val = clean(m.group("val"))
                if not val and j + 1 < len(row):
                    val = clean(row[j + 1]["t"])
                if not val and i + 1 < len(rows):
                    w = c["x1"] - c["x0"]
                    below = [d for d in rows[i + 1] if d["x1"] >= c["x0"] - w and d["x0"] <= c["x1"] + w]
                    if below:
                        val = clean(below[0]["t"])
                return val, i
    return "", None


def token_with_digit(s):
    m = re.search(r"[A-Z0-9][A-Z0-9\-/]*\d[A-Z0-9\-/]*", str(s), re.I)
    return m.group(0) if m else ""


MONTHS = "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec"
DATE_RE = re.compile(
    r"(\d{4}[-/.]\d{1,2}[-/.]\d{1,2}|\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4}|"
    r"(?:%s)[a-z]*\.?\s*\d{1,2},?\s*\d{4}|\d{1,2}[\s-]*(?:%s)[a-z]*[\s,.-]*\d{2,4})" % (MONTHS, MONTHS),
    re.I)


def parse_date(s, dayfirst=False):
    s = re.sub(r"([A-Za-z])(\d)", r"\1 \2", s)   # "Oct21,2026" -> "Oct 21, 2026"
    s = re.sub(r",\s*", ", ", s)
    try:
        yearfirst = bool(re.match(r"\d{4}", s))
        d = pd.to_datetime(s, dayfirst=dayfirst and not yearfirst, yearfirst=yearfirst)
        return d.strftime("%Y-%m-%d") if 2000 <= d.year <= 2100 else ""
    except Exception:
        return ""


def first_date(text, dayfirst=False):
    for m in DATE_RE.finditer(str(text)):
        d = parse_date(m.group(0), dayfirst)
        if d:
            return d
    return ""


def labeled_date(rows, labels, dayfirst):
    val, i = find_value(rows, labels)
    if i is None:
        return ""
    return first_date(val, dayfirst) or first_date(rt(rows[i]), dayfirst) or \
        (first_date(rt(rows[i + 1]), dayfirst) if i + 1 < len(rows) else "")


def to_num(v):
    try:
        if v in (None, ""):
            return None
        s = re.sub(r"[^\d.\-]", "", str(v).replace(",", ""))
        return float(s) if s not in ("", ".", "-") else None
    except ValueError:
        return None


NUM_RE = re.compile(r"^\$?\d{1,3}(,\d{3})*(\.\d+)?$|^\$?\d+(\.\d+)?$")


def is_num(t):
    return bool(NUM_RE.match(t.strip()))


def is_money(t):
    t = t.strip()
    return is_num(t) and ("$" in t or re.search(r"\.\d{2}$", t) is not None)


def auto_category(desc):
    d = str(desc).lower()
    for cat, words in CATEGORY_KEYWORDS.items():
        if any(w in d for w in words):
            return cat
    return "Other"


def company_line(rows, exclude=r"bird"):
    """Company name near the top (Ltd/Inc/Corp/Supply...), not Bird."""
    suffix = re.compile(r"(ltd|limited|inc|corp|corporation|llc|ulc)\.?(\b|$)|"
                        r"\b(supply|supplies|group|industries|rentals?|contracting|contractors?)\b", re.I)
    for row in rows[:12]:
        for c in row:
            if suffix.search(c["t"]) and not re.search(exclude, c["t"], re.I):
                return clean(c["t"])
    for row in rows[:5]:
        for c in row:
            if len(re.findall(r"[A-Za-z]", c["t"])) >= 4 and not re.search(
                    exclude + r"|packing|invoice|slip|ticket|page|date|time\s*&|t\s*&\s*m", c["t"], re.I):
                return clean(c["t"])
    return ""


# ---------------- table alignment ----------------
def map_header(row, rules):
    """rules: ordered [(field, regex)]. Returns [(x0, x1, field)] or None."""
    hdr, fields = [], set()
    for c in row:
        if ":" in c["t"] or len(c["t"].split()) > 4:   # "Ship To:" etc. are labels, not headers
            continue
        for field, rx in rules:
            if re.search(rx, c["t"], re.I):
                hdr.append((c["x0"], c["x1"], field))
                fields.add(field)
                break
    digits = sum(ch.isdigit() for ch in rt(row))
    ok = len(fields) >= 2 and digits <= 4 and len(hdr) * 2 >= len(row)
    return hdr if ok else None


def align(row, hdr):
    out = {}
    for c in row:
        best = max(hdr, key=lambda h: min(c["x1"], h[1]) - max(c["x0"], h[0]))
        if min(c["x1"], best[1]) - max(c["x0"], best[0]) <= 0:
            cx = (c["x0"] + c["x1"]) / 2
            best = min(hdr, key=lambda h: abs(cx - (h[0] + h[1]) / 2))
        out[best[2]] = (out.get(best[2], "") + " " + c["t"]).strip()
    return out


def split_qty_uom(s):
    toks = str(s).split()
    qty = next((to_num(t) for t in toks if is_num(t)), None)
    uom = next((t.upper() for t in toks if t.upper().strip(".") in UOMS), "")
    return qty, uom


# ============================================================================
# 3. MATERIAL INVENTORY PARSER
# ============================================================================
INV_HDR = [
    ("Notes", r"b\s*/\s*o\b|back\s*-?\s*order|comment|note"),
    ("Unit Price", r"unit\s*(price|cost)|price|rate"),
    ("Extension", r"\bext|amount|line\s*total|^total"),
    ("U/M", r"u\s*/\s*m|uom|\bum\b|^unit$|units?\s*of"),
    ("Description", r"desc|material"),
    ("Product #", r"item|part|product|sku|code|cat\s*(#|no)"),
    ("_Ordered", r"order|\bord\b"),
    ("Qty", r"qty|quant|ship|shp|rec'?d|received|deliv"),
]
INV_STOP = re.compile(r"sub\s*-?\s*total|^\s*total|\btax\b|\bgst\b|\bpst\b|\bhst\b|thank|"
                      r"signature|received\s*by|receiver|page\s*\d|freight|terms|balance|"
                      r"amount\s*due|total\s*(qty|weight|pieces)", re.I)
ADDRESS = re.compile(r"\b(st|street|ave|avenue|rd|road|blvd|drive|dr|nw|ne|sw|se)\b\.?\s|"
                     r"\b(ab|alberta|bc|sk|on)\b\s*[A-Z]\d[A-Z]|phone|tel|fax|www\.|@|ship\s*to|"
                     r"bill\s*to|sold\s*to", re.I)


def doc_type(text):
    t = re.sub(r"\s", "", text.lower())
    for key, name in (("packingslip", "Packing Slip"), ("packinglist", "Packing Slip"),
                      ("deliveryticket", "Delivery Ticket"), ("deliveryslip", "Delivery Ticket"),
                      ("deliveryreceipt", "Delivery Ticket"), ("billoflading", "Bill of Lading"),
                      ("invoice", "Invoice"), ("quotation", "Quote"), ("quote", "Quote"),
                      ("receipt", "Receipt")):
        if key in t:
            return name
    return "Other"


def parse_item_tokens(text):
    """Fallback when no table header was found."""
    toks = text.split()
    product = ""
    if toks and re.search(r"\d", toks[0]) and (re.search(r"[A-Za-z\-]", toks[0]) or len(toks[0]) >= 4) \
            and sum(is_num(t) for t in toks[1:]) >= 1:
        product = toks.pop(0)
    uom = next((t.upper() for t in toks if t.upper().strip(".") in UOMS), "")
    nums = [t for t in toks if is_num(t)]
    desc = " ".join(t for t in toks if not is_num(t) and t.upper().strip(".") != uom)
    qty = price = ext = None
    if len(nums) >= 3:
        qty, price, ext = to_num(nums[0]), to_num(nums[-2]), to_num(nums[-1])
    elif len(nums) == 2:
        if is_money(nums[1]):
            qty, price = to_num(nums[0]), to_num(nums[1])
        else:
            qty = to_num(nums[1])
    elif len(nums) == 1:
        qty = to_num(nums[0])
    return {"Product #": product, "Description": desc, "Qty": qty, "U/M": uom,
            "Unit Price": price, "Extension": ext}


def parse_inventory(rows, filename, received, ordered_by_default, dayfirst):
    text = "\n".join(rt(r) for r in rows)
    vendor, _ = find_value(rows, [r"vendor", r"supplier", r"sold\s*by", r"shipped\s*from", r"remit\s*to"])
    if not vendor or re.search(r"bird", vendor, re.I):
        vendor = company_line(rows)
    doc_no, _ = find_value(rows, [r"packing\s*slip\s*(#|no\.?|number)", r"invoice\s*(#|no\.?|number)",
                                  r"delivery\s*(ticket|slip)?\s*(#|no\.?|number)", r"ticket\s*(#|no\.?)",
                                  r"(sales\s*)?order\s*(#|no\.?|number)", r"document\s*(#|no\.?)",
                                  r"b\s*/?\s*o?\s*l\s*(#|no\.?)", r"shipment\s*(#|no\.?)", r"slip\s*(#|no\.?)"])
    po, _ = find_value(rows, [r"(customer\s*)?p\.?\s*o\.?\s*(#|no\.?|number|num)",
                              r"purchase\s*order\s*(#|no\.?|number)?", r"\bp\.?\s*o\b"])
    ordered_by, _ = find_value(rows, [r"order(ed)?\s*by", r"requested\s*by", r"attn|attention",
                                      r"buyer", r"contact"])
    doc_date = labeled_date(rows, [r"(invoice|ship|order|document|delivery)?\s*date"], dayfirst) \
        or first_date(text, dayfirst)

    header = {
        "Doc Date": doc_date,
        "Date Received": received.strftime("%Y-%m-%d"),
        "Vendor": vendor,
        "Doc Type": doc_type(text),
        "Doc #": token_with_digit(doc_no),
        "PO #": token_with_digit(po),
        "Ordered By": clean(re.sub(DATE_RE, "", ordered_by)) or ordered_by_default,
        "Source File": filename,
    }

    hdr, start = None, 0
    for i, row in enumerate(rows):
        hdr = map_header(row, INV_HDR)
        if hdr:
            start = i + 1
            break
    if hdr and not any(f == "Qty" for *_, f in hdr):
        hdr = [(a, b, "Qty" if f == "_Ordered" else f) for a, b, f in hdr]

    items = []
    for row in rows[start:]:
        line = rt(row)
        if INV_STOP.search(line):
            if hdr:
                break
            continue
        if len(re.findall(r"[A-Za-z]", line)) < 3 or not re.search(r"\d", line):
            continue
        if not hdr and (ADDRESS.search(line) or DATE_RE.search(line) or re.search(r"#|no\.", line, re.I)):
            continue
        if hdr:
            f = align(row, hdr)
            qty, uom = split_qty_uom(f.get("Qty", ""))
            if not f.get("U/M") and uom:
                f["U/M"] = uom
            notes = f.get("Notes", "")
            item = {"Product #": f.get("Product #", ""), "Description": f.get("Description", ""),
                    "Qty": qty, "U/M": (f.get("U/M") or "").upper(),
                    "Unit Price": to_num(f.get("Unit Price")), "Extension": to_num(f.get("Extension")),
                    "Notes": ("B/O: " + notes) if notes and is_num(notes) else notes}
            if not item["Description"]:
                item["Description"] = item["Product #"]
            if qty is None and not item["Description"]:
                continue
        else:
            item = parse_item_tokens(line)
            item["Notes"] = ""
        if item["Extension"] is None and item["Qty"] is not None and item["Unit Price"] is not None:
            item["Extension"] = round(item["Qty"] * item["Unit Price"], 2)
        item["Category"] = auto_category(item["Description"])
        items.append({**header, **item})

    if not items:
        items.append({**header, "Notes": "No line items read – enter manually from photo"})
    return items, text


# ============================================================================
# 4. TIME & MATERIAL PARSER
# ============================================================================
SEC_LAB = re.compile(r"\blabou?r\b|manpower|\bworkers?\b|\bemployees?\b|\bcrew\b|"
                     r"\bname\b.*\b(hrs|hours|class)", re.I)
SEC_MAT = re.compile(r"\bmaterials?\b", re.I)
SEC_EQP = re.compile(r"\bequipment\b|\bequip\b|\brentals?\b|\btools?\b", re.I)
SEC_END = re.compile(r"^\s*(grand\s*)?total|signature|signed|approved|authori[sz]ed|accepted|"
                     r"comments|remarks|description\s*of\s*work|scope|submitted\s*by", re.I)
LAB_HDR = [
    ("Total Hrs", r"total\s*h(ou)?rs?|total\s*hours"),
    ("OT Hrs", r"\bo\.?\s*/?\s*t\b|overtime|1\.5|double|\bdt\b"),
    ("Reg Hrs", r"\breg|\bst\b|straight|hours|hrs|\bhr\b"),
    ("Unit Rate", r"rate"),
    ("Amount", r"amount|total|cost|\$"),
    ("Classification", r"class|trade|position|title|craft|occupation"),
    ("Worker Name", r"name|employee|worker|labou?r|crew"),
]
MAT_HDR = [
    ("Unit Rate", r"unit\s*(price|cost)|rate|price"),
    ("Amount", r"amount|total|\bext|cost"),
    ("U/M", r"u\s*/\s*m|uom|^unit$"),
    ("Qty", r"qty|quant|hrs|hours|days|^#$"),
    ("Material / Equipment", r"desc|material|item|equip|type|rental"),
]
TRADE_SIG = [r"contractor'?s?\s*(signature|rep(resentative)?|foreman|name)",
             r"sub\s*-?\s*contractor\s*(signature|rep)", r"trade\s*(signature|rep)",
             r"submitted\s*by", r"prepared\s*by", r"foreman\s*(signature|name)?", r"signed\s*by"]
BIRD_SIG = [r"approved\s*by", r"authori[sz]ed\s*by", r"accepted\s*by", r"verified\s*by",
            r"bird\s*(rep|representative|signature|approval|super\w*|foreman)",
            r"general\s*contractor", r"superintendent", r"site\s*(rep|supervisor)",
            r"\bg\.?\s*c\.?\s*(signature|rep)"]


def sig_name(rows, labels, dayfirst):
    val, i = find_value(rows, labels)
    if i is None:
        return "NOT FOUND – verify", ""
    d = first_date(rt(rows[i]), dayfirst) or (first_date(rt(rows[i + 1]), dayfirst) if i + 1 < len(rows) else "")
    name = re.sub(DATE_RE, "", val)
    name = re.sub(r"(?i)signature|sign(ed)?|print(ed)?|name|date|x\b", "", name)
    name = clean(name)
    if len(re.findall(r"[A-Za-z]", name)) >= 2:
        return name.title(), d
    return "Signature line – verify", d


def parse_labour(row, hdr):
    if hdr:
        f = align(row, hdr)
        reg, ot = to_num(f.get("Reg Hrs")), to_num(f.get("OT Hrs"))
        tot = to_num(f.get("Total Hrs"))
        if reg is None and tot is not None:
            reg = tot - (ot or 0)
        name = f.get("Worker Name", "")
        cls = f.get("Classification", "")
        m = re.search(CLASSIFICATIONS, name, re.I)
        if m and not cls:
            cls, name = m.group(0), clean(name.replace(m.group(0), ""))
        return {"Worker Name": clean(name).title() or "Unnamed", "Classification": cls.title(),
                "Reg Hrs": reg or 0, "OT Hrs": ot or 0,
                "Unit Rate": to_num(f.get("Unit Rate")), "Amount": to_num(f.get("Amount"))}
    text = rt(row)
    m = re.search(CLASSIFICATIONS, text, re.I)
    cls = m.group(0) if m else ""
    rest = text.replace(cls, " ") if cls else text
    toks = rest.split()
    hrs = [to_num(t) for t in toks if is_num(t) and not is_money(t) and (to_num(t) or 0) <= 24]
    money = [to_num(t) for t in toks if is_money(t) or (is_num(t) and (to_num(t) or 0) > 24)]
    name = " ".join(t for t in toks if not is_num(t) and re.search(r"[A-Za-z]", t))
    return {"Worker Name": clean(name).title() or "Unnamed", "Classification": cls.title(),
            "Reg Hrs": hrs[0] if hrs else 0, "OT Hrs": hrs[1] if len(hrs) > 1 else 0,
            "Unit Rate": money[0] if money else None,
            "Amount": money[-1] if len(money) > 1 else None}


def parse_matequip(row, hdr):
    if hdr:
        f = align(row, hdr)
        qty, uom = split_qty_uom(f.get("Qty", ""))
        return {"Material / Equipment": f.get("Material / Equipment", ""), "Qty": qty,
                "U/M": (f.get("U/M") or uom).upper(), "Unit Rate": to_num(f.get("Unit Rate")),
                "Amount": to_num(f.get("Amount"))}
    it = parse_item_tokens(rt(row))
    return {"Material / Equipment": (it["Product #"] + " " + it["Description"]).strip(),
            "Qty": it["Qty"], "U/M": it["U/M"], "Unit Rate": it["Unit Price"], "Amount": it["Extension"]}


def scope_text(rows):
    val, i = find_value(rows, [r"description\s*of\s*(the\s*)?(extra\s*)?work", r"work\s*description",
                               r"scope\s*(of\s*work)?", r"work\s*performed", r"extra\s*work",
                               r"description", r"details\s*of\s*work", r"reason"])
    if i is None:
        return ""
    parts = [val] if val else []
    for row in rows[i + 1:i + 5]:
        line = rt(row)
        if SEC_LAB.search(line) or SEC_MAT.search(line) or SEC_EQP.search(line) or \
                SEC_END.search(line) or re.search(r"\w+\s*:", line):
            break
        parts.append(line)
    return clean(" ".join(dict.fromkeys(parts)))


def parse_tm(rows, filename, dayfirst):
    text = "\n".join(rt(r) for r in rows)
    company, _ = find_value(rows, [r"sub\s*-?\s*contractor(?!'?s?\s*(sig|rep))",
                                   r"contractor(?!'?s?\s*(sig|rep|foreman))",
                                   r"company", r"trade(?!\s*(sig|rep))", r"vendor", r"\bfrom\b"])
    if not company or re.search(r"bird", company, re.I):
        company = company_line(rows)
    ticket, _ = find_value(rows, [r"t\s*&\s*m\s*(ticket)?\s*(#|no\.?|number)", r"ticket\s*(#|no\.?|number)",
                                  r"sheet\s*(#|no\.?)", r"work\s*order\s*(#|no\.?)", r"\bno\.", r"#"])
    work_date = labeled_date(rows, [r"date\s*of\s*work", r"work\s*date", r"date\s*worked",
                                    r"date\s*performed"], dayfirst)
    ticket_date = labeled_date(rows, [r"(ticket\s*)?date(?!\s*signed)"], dayfirst) or first_date(text, dayfirst)
    refs = sorted(set(m.group(0).upper().replace(" ", "") for m in re.finditer(
        r"\b(RFI|SI|CD|CO|CCN|PCN|SWI|CCO|PCO)\s*[-#:]?\s*\d+[A-Z0-9\-]*", text, re.I)))
    location, _ = find_value(rows, [r"location", r"area", r"level", r"grid\s*(lines?)?"])
    trade_sig, d1 = sig_name(rows, TRADE_SIG, dayfirst)
    bird_sig, d2 = sig_name(rows, BIRD_SIG, dayfirst)
    signed_date = labeled_date(rows, [r"date\s*signed", r"sign(ed)?\s*date"], dayfirst) or d2 or d1

    base = {
        "Work Date": work_date or ticket_date, "Ticket Date": ticket_date,
        "Trade Company": company, "T&M Ticket #": token_with_digit(ticket),
        "Scope of Work": scope_text(rows), "Reference (RFI/SI/CD)": ", ".join(refs),
        "Location": location, "Trade Signed By": trade_sig, "Approved By (Bird)": bird_sig,
        "Signed Date": signed_date, "Source File": filename,
    }

    out, section, hdr = [], None, None
    for row in rows:
        line = rt(row)
        digits = sum(ch.isdigit() for ch in line)
        new_sec = None
        if digits <= 4:
            if SEC_EQP.search(line):
                new_sec = "Equipment"
            elif SEC_MAT.search(line):
                new_sec = "Material"
            elif SEC_LAB.search(line):
                new_sec = "Labour"
        if new_sec:
            section = new_sec
            hdr = map_header(row, LAB_HDR if section == "Labour" else MAT_HDR)
            continue
        if section is None:
            continue
        if SEC_END.search(line):
            section, hdr = None, None
            continue
        if hdr is None:
            h = map_header(row, LAB_HDR if section == "Labour" else MAT_HDR)
            if h:
                hdr = h
                continue
        if not re.search(r"\d", line) or len(re.findall(r"[A-Za-z]", line)) < 2:
            continue
        if section == "Labour":
            r = parse_labour(row, hdr)
            reg, ot = r["Reg Hrs"] or 0, r["OT Hrs"] or 0
            amt = r["Amount"]
            if amt is None and r["Unit Rate"] is not None:
                amt = round((reg + ot) * r["Unit Rate"], 2)
            out.append({**base, "Line Type": "Labour", "Worker Name": r["Worker Name"],
                        "Classification": r["Classification"], "Reg Hrs": reg, "OT Hrs": ot,
                        "Total Hrs": reg + ot, "Material / Equipment": "", "Qty": reg + ot,
                        "U/M": "HR", "Unit Rate": r["Unit Rate"], "Amount": amt, "Notes": ""})
        else:
            r = parse_matequip(row, hdr)
            amt = r["Amount"]
            if amt is None and r["Qty"] is not None and r["Unit Rate"] is not None:
                amt = round(r["Qty"] * r["Unit Rate"], 2)
            if section == "Equipment" and not r["U/M"]:
                r["U/M"] = "HR"
            is_hr = str(r["U/M"]).upper().startswith("H")
            out.append({**base, "Line Type": section, "Worker Name": "", "Classification": "",
                        "Reg Hrs": None, "OT Hrs": None,
                        "Total Hrs": r["Qty"] if section == "Equipment" and is_hr else None,
                        "Material / Equipment": r["Material / Equipment"], "Qty": r["Qty"],
                        "U/M": r["U/M"], "Unit Rate": r["Unit Rate"], "Amount": amt, "Notes": ""})
    if not out:
        out.append({**base, "Line Type": "", "Notes": "No labour/material lines read – enter manually from photo"})
    return out, text


# ============================================================================
# 5. EXCEL BUILDERS
# ============================================================================
NAVY, LIGHT = "003B71", "EAF0F7"
THIN = Side(style="thin", color="B7C4D4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '"$"#,##0.00'
HDR = 4


def needs_sig(series):
    s = series.astype(str).str.upper()
    return s.str.contains("NOT ") | s.str.contains("VERIFY") | (s.str.strip() == "")


def style_header(cell):
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def title(ws, text):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = f"Generated: {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = Font(italic=True, size=9, color="666666")


def as_date(v):
    if v in (None, "") or (not isinstance(v, str) and pd.isna(v)):
        return None
    try:
        return pd.to_datetime(v).date()
    except Exception:
        return v


def write_log(ws, df, cols, widths, money, nums, dates, wrap, formulas=None):
    cmap = {n: i + 1 for i, n in enumerate(cols)}
    for n, c in cmap.items():
        style_header(ws.cell(row=HDR, column=c, value=n))
    r = HDR
    for r, (_, row) in enumerate(df.reindex(columns=cols).iterrows(), HDR + 1):
        for n in cols:
            v = row[n]
            v = None if (not isinstance(v, str) and pd.isna(v)) else v
            if n in dates:
                v = as_date(v)
            if formulas and n in formulas:
                v = formulas[n](r, cmap, row)
            cell = ws.cell(row=r, column=cmap[n], value=v)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=n in wrap)
            if (r - HDR) % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT)
            if n in money:
                cell.number_format = MONEY
            elif n in nums:
                cell.number_format = "#,##0.##"
            elif n in dates:
                cell.number_format = "yyyy-mm-dd"
    for n, w in widths.items():
        ws.column_dimensions[get_column_letter(cmap[n])].width = w
    ws.row_dimensions[HDR].height = 32
    ws.freeze_panes = ws.cell(row=HDR + 1, column=1)
    last = max(r, HDR + 1)
    ws.auto_filter.ref = f"A{HDR}:{get_column_letter(len(cols))}{last}"
    return cmap, last


def add_dropdown(ws, wb, col_letter, values, last, list_col):
    lists = wb["Lists"] if "Lists" in wb.sheetnames else wb.create_sheet("Lists")
    for i, v in enumerate(values, 2):
        lists.cell(row=i, column=list_col, value=v)
    L = get_column_letter(list_col)
    dv = DataValidation(type="list", formula1=f"=Lists!${L}$2:${L}${len(values) + 1}", allow_blank=True)
    dv.add(f"{col_letter}{HDR + 1}:{col_letter}{last + 200}")
    ws.add_data_validation(dv)


def summary_table(ws, start_row, heading, headers, rows, money_idx=(), num_idx=()):
    ws.cell(row=start_row, column=1, value=heading).font = Font(bold=True, size=12, color=NAVY)
    for c, h in enumerate(headers, 1):
        style_header(ws.cell(row=start_row + 1, column=c, value=h))
    r = start_row + 1
    for r, vals in enumerate(rows, start_row + 2):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            if c - 1 in money_idx:
                cell.number_format = MONEY
            elif c - 1 in num_idx:
                cell.number_format = "#,##0.##"
    return r + 3


def inv_sheets(wb, df, project):
    ws = wb.create_sheet("Material Inventory")
    title(ws, f"{project} – Material Inventory Log")

    def ext_formula(r, cm, row):
        q, p = get_column_letter(cm["Qty"]), get_column_letter(cm["Unit Price"])
        fb = row["Extension"] if pd.notna(row["Extension"]) and row["Extension"] != "" else '""'
        return f"=IF(AND(ISNUMBER({q}{r}),ISNUMBER({p}{r})),{q}{r}*{p}{r},{fb})"

    widths = {"Doc Date": 12, "Date Received": 13, "Vendor": 24, "Doc Type": 15, "Doc #": 14,
              "PO #": 14, "Ordered By": 16, "Product #": 16, "Description": 42, "Category": 20,
              "Qty": 8, "U/M": 7, "Unit Price": 12, "Extension": 13, "Notes": 32}
    cm, last = write_log(ws, df, INV_COLUMNS, widths, {"Unit Price", "Extension"}, {"Qty"},
                         {"Doc Date", "Date Received"}, {"Description", "Notes"},
                         {"Extension": ext_formula})
    ext = get_column_letter(cm["Extension"])
    ws.cell(row=last + 1, column=cm["Unit Price"], value="TOTAL").font = Font(bold=True)
    t = ws.cell(row=last + 1, column=cm["Extension"], value=f"=SUM({ext}{HDR + 1}:{ext}{last})")
    t.font, t.number_format = Font(bold=True), MONEY
    add_dropdown(ws, wb, get_column_letter(cm["Doc Type"]), DOC_TYPES, last, 1)
    add_dropdown(ws, wb, get_column_letter(cm["Category"]), CATEGORIES, last, 2)

    s = wb.create_sheet("Inventory Summary")
    title(s, f"{project} – Material Inventory Summary")
    sh = "'Material Inventory'"
    rng = lambda n: f"{sh}!${get_column_letter(cm[n])}${HDR + 1}:${get_column_letter(cm[n])}${last}"
    row = summary_table(s, 4, "Overall", ["Metric", "Value"], [
        ["Line items", f"=COUNTA({rng('Description')})"],
        ["Documents", df["Source File"].replace("", pd.NA).nunique()],
        ["Vendors", df["Vendor"].replace("", pd.NA).nunique()],
        ["Total value", f"=SUM({rng('Extension')})"],
    ])
    s.cell(row=row - 3, column=2).number_format = MONEY
    vendors = sorted(v for v in df["Vendor"].dropna().unique() if v)
    row = summary_table(s, row, "By Vendor", ["Vendor", "Documents", "Line Items", "Total Value"],
                        [[v, df.loc[df["Vendor"] == v, "Source File"].nunique(),
                          f'=COUNTIF({rng("Vendor")},A{row + 2 + i})',
                          f'=SUMIF({rng("Vendor")},A{row + 2 + i},{rng("Extension")})']
                         for i, v in enumerate(vendors)], money_idx=(3,))
    cats = sorted(c for c in df["Category"].dropna().unique() if c)
    row = summary_table(s, row, "By Category", ["Category", "Line Items", "Total Qty", "Total Value"],
                        [[c, f'=COUNTIF({rng("Category")},A{row + 2 + i})',
                          f'=SUMIF({rng("Category")},A{row + 2 + i},{rng("Qty")})',
                          f'=SUMIF({rng("Category")},A{row + 2 + i},{rng("Extension")})']
                         for i, c in enumerate(cats)], money_idx=(3,), num_idx=(2,))
    docs = df.drop_duplicates(["Source File"])
    end = summary_table(s, row, "By Document",
                        ["Doc Date", "Date Received", "Vendor", "Doc Type", "Doc #", "PO #", "Line Items", "Value"],
                        [[as_date(d["Doc Date"]), as_date(d["Date Received"]), d["Vendor"], d["Doc Type"],
                          d["Doc #"], d["PO #"], int((df["Source File"] == d["Source File"]).sum()),
                          pd.to_numeric(df.loc[df["Source File"] == d["Source File"], "Extension"],
                                        errors="coerce").sum()]
                         for _, d in docs.iterrows()], money_idx=(7,))
    for r in range(row + 2, end - 2):
        for c in (1, 2):
            s.cell(row=r, column=c).number_format = "yyyy-mm-dd"
    for col, w in zip("ABCDEFGH", (30, 14, 24, 16, 14, 14, 12, 14)):
        s.column_dimensions[col].width = w


def tm_sheets(wb, df, project):
    ws = wb.create_sheet("T&M Log")
    title(ws, f"{project} – Time & Material Tracking Log")

    def hrs_formula(r, cm, row):
        a, b = get_column_letter(cm["Reg Hrs"]), get_column_letter(cm["OT Hrs"])
        fb = row["Total Hrs"] if pd.notna(row["Total Hrs"]) and row["Total Hrs"] != "" else '""'
        return f'=IF(COUNT({a}{r},{b}{r})=0,{fb},N({a}{r})+N({b}{r}))'

    widths = {"Work Date": 12, "Ticket Date": 12, "Trade Company": 24, "T&M Ticket #": 13,
              "Scope of Work": 40, "Reference (RFI/SI/CD)": 14, "Location": 16, "Line Type": 11,
              "Worker Name": 20, "Classification": 15, "Reg Hrs": 8, "OT Hrs": 8, "Total Hrs": 9,
              "Material / Equipment": 32, "Qty": 8, "U/M": 7, "Unit Rate": 11, "Amount": 12,
              "Trade Signed By": 22, "Approved By (Bird)": 22, "Signed Date": 12, "Notes": 30}
    cm, last = write_log(ws, df, TM_COLUMNS, widths, {"Unit Rate", "Amount"},
                         {"Reg Hrs", "OT Hrs", "Total Hrs", "Qty"},
                         {"Work Date", "Ticket Date", "Signed Date"},
                         {"Scope of Work", "Material / Equipment", "Notes"},
                         {"Total Hrs": hrs_formula})
    red = PatternFill("solid", fgColor="F8D7DA")
    for r in range(HDR + 1, last + 1):
        for n in ("Trade Signed By", "Approved By (Bird)"):
            c = ws.cell(row=r, column=cm[n])
            v = str(c.value or "").upper()
            if "NOT " in v or "VERIFY" in v:
                c.fill, c.font = red, Font(bold=True, color="9C0006")
    for n in ("Total Hrs", "Amount"):
        L = get_column_letter(cm[n])
        c = ws.cell(row=last + 1, column=cm[n], value=f"=SUM({L}{HDR + 1}:{L}{last})")
        c.font = Font(bold=True)
        c.number_format = MONEY if n == "Amount" else "#,##0.##"
    ws.cell(row=last + 1, column=cm["OT Hrs"], value="TOTAL").font = Font(bold=True)
    add_dropdown(ws, wb, get_column_letter(cm["Line Type"]), LINE_TYPES, last, 3)

    s = wb.create_sheet("T&M Summary")
    title(s, f"{project} – Time & Material Summary")
    sh = "'T&M Log'"
    rng = lambda n: f"{sh}!${get_column_letter(cm[n])}${HDR + 1}:${get_column_letter(cm[n])}${last}"
    tickets = df.drop_duplicates(["Source File"])
    unsigned = tickets[needs_sig(tickets["Trade Signed By"]) | needs_sig(tickets["Approved By (Bird)"])]
    row = summary_table(s, 4, "Overall", ["Metric", "Value"], [
        ["T&M tickets", len(tickets)],
        ["Trade companies", df["Trade Company"].replace("", pd.NA).nunique()],
        ["Labour hours", f'=SUMIF({rng("Line Type")},"Labour",{rng("Total Hrs")})'],
        ["Equipment hours", f'=SUMIF({rng("Line Type")},"Equipment",{rng("Total Hrs")})'],
        ["Total amount", f"=SUM({rng('Amount')})"],
        ["Tickets to check for signature", len(unsigned)],
    ])
    s.cell(row=row - 4, column=2).number_format = MONEY

    comps = sorted(c for c in df["Trade Company"].dropna().unique() if c)
    rows = []
    for i, c in enumerate(comps):
        a = f"A{row + 2 + i}"
        sub = df[df["Trade Company"] == c]
        rows.append([
            c, sub["Source File"].nunique(),
            sub.loc[sub["Line Type"] == "Labour", "Worker Name"].nunique(),
            f'=SUMIFS({rng("Total Hrs")},{rng("Trade Company")},{a},{rng("Line Type")},"Labour")',
            f'=SUMIFS({rng("Amount")},{rng("Trade Company")},{a},{rng("Line Type")},"Labour")',
            f'=SUMIFS({rng("Amount")},{rng("Trade Company")},{a},{rng("Line Type")},"Material")',
            f'=SUMIFS({rng("Amount")},{rng("Trade Company")},{a},{rng("Line Type")},"Equipment")',
            f'=SUMIF({rng("Trade Company")},{a},{rng("Amount")})',
        ])
    row = summary_table(s, row, "By Trade Company",
                        ["Trade Company", "Tickets", "Workers", "Labour Hrs", "Labour $",
                         "Material $", "Equipment $", "Total $"],
                        rows, money_idx=(4, 5, 6, 7), num_idx=(3,))

    trows = []
    for _, t in tickets.sort_values(["Work Date", "Trade Company"]).iterrows():
        sub = df[df["Source File"] == t["Source File"]]
        lab = sub[sub["Line Type"] == "Labour"]
        mats = "; ".join(x for x in sub.loc[sub["Line Type"] == "Material", "Material / Equipment"].astype(str) if x)
        eqp = "; ".join(x for x in sub.loc[sub["Line Type"] == "Equipment", "Material / Equipment"].astype(str) if x)
        ok = not (needs_sig(pd.Series([t["Trade Signed By"]])).iloc[0]
                  or needs_sig(pd.Series([t["Approved By (Bird)"]])).iloc[0])
        trows.append([as_date(t["Work Date"]), t["Trade Company"], t["T&M Ticket #"], t["Scope of Work"],
                      ", ".join(x for x in lab["Worker Name"].astype(str).unique() if x),
                      pd.to_numeric(lab["Total Hrs"], errors="coerce").sum(), mats, eqp,
                      pd.to_numeric(sub["Amount"], errors="coerce").sum(),
                      t["Trade Signed By"], t["Approved By (Bird)"], as_date(t["Signed Date"]),
                      "✔ Signed" if ok else "⚠ Check signature"])
    end = summary_table(s, row, "By Ticket",
                        ["Work Date", "Trade Company", "Ticket #", "Scope", "Workers", "Labour Hrs",
                         "Materials", "Equipment", "Total $", "Trade Signed By", "Approved By (Bird)",
                         "Signed Date", "Status"],
                        trows, money_idx=(8,), num_idx=(5,))
    for r in range(row + 2, end - 2):
        for c in (1, 12):
            s.cell(row=r, column=c).number_format = "yyyy-mm-dd"
        for c in (4, 5, 7, 8):
            s.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHIJKLM", (24, 22, 12, 40, 28, 12, 30, 26, 14, 22, 22, 12, 18)):
        s.column_dimensions[col].width = w


def make_workbook(project, inv=None, tm=None, texts=None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    if inv is not None:
        inv_sheets(wb, inv, project)
    if tm is not None:
        tm_sheets(wb, tm, project)
    if "Lists" in wb.sheetnames:
        wb.move_sheet("Lists", offset=len(wb.sheetnames))
    src = wb.create_sheet("Source Text")
    src.append(["Module", "Source File", "Rows Extracted", "Text Read From Photo/PDF"])
    for c in src[1]:
        style_header(c)
    for name, df in (("Material Inventory", inv), ("T&M", tm)):
        if df is not None and not df.empty:
            for f, n in df.groupby("Source File").size().items():
                src.append([name, f, int(n), (texts or {}).get(f, "")[:32000]])
                src.cell(row=src.max_row, column=4).alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCD", (20, 34, 14, 100)):
        src.column_dimensions[col].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def clean_inv(df):
    df = df.copy()
    for c in ("Qty", "Unit Price", "Extension"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    calc = (df["Qty"] * df["Unit Price"]).round(2)
    df["Extension"] = calc.where(calc.notna(), df["Extension"])
    return df.fillna({c: "" for c in df.columns if df[c].dtype == object})


def clean_tm(df):
    df = df.copy()
    for c in ("Reg Hrs", "OT Hrs", "Total Hrs", "Qty", "Unit Rate", "Amount"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    has = df["Reg Hrs"].notna() | df["OT Hrs"].notna()
    df["Total Hrs"] = (df["Reg Hrs"].fillna(0) + df["OT Hrs"].fillna(0)).where(has, df["Total Hrs"])
    return df.fillna({c: "" for c in df.columns if df[c].dtype == object})


# ============================================================================
# 6. SIDEBAR  (no API key / endpoint fields)
# ============================================================================
with st.sidebar:
    st.header("⚙️ Settings")
    project = st.text_input("Project", "MMCC – McMurray Métis Cultural Centre")
    date_order = st.radio("Dates written like 03/04/2026 mean", ["Month/Day (MM/DD)", "Day/Month (DD/MM)"])
    dayfirst = date_order.startswith("Day")
    st.success("✅ Runs offline – no API key, no endpoint, no extra software.")
    if not HEIC_OK:
        st.warning("HEIC not supported – run `pip install pillow-heif`")
    if not PDF_OK:
        st.warning("PDF not supported – run `pip install pymupdf`")
    st.divider()
    st.download_button("⬇️ Blank Templates (Excel)",
                       make_workbook(project, pd.DataFrame(columns=INV_COLUMNS + ["Source File"]),
                                     pd.DataFrame(columns=TM_COLUMNS + ["Source File"])),
                       "MMCC_Inventory_TM_Template.xlsx", width="stretch")
    if st.button("🗑️ Clear all data", width="stretch"):
        for k, cols in (("inv", INV_COLUMNS), ("tm", TM_COLUMNS)):
            st.session_state[f"{k}_base"] = pd.DataFrame(columns=cols + ["Source File"])
            st.session_state[f"{k}_view"] = st.session_state[f"{k}_base"]
            st.session_state[f"{k}_processed"] = set()
            st.session_state[f"{k}_images"] = {}
            st.session_state[f"{k}_text"] = {}
            st.session_state[f"{k}_ver"] += 1
        st.rerun()


# ============================================================================
# 7. SHARED UI
# ============================================================================
def upload_and_extract(mod, parse_fn, label):
    files = st.file_uploader(f"Upload {label} – HEIC / JPG / PNG / WEBP / PDF (multiple allowed)",
                             type=FILE_TYPES, accept_multiple_files=True, key=f"{mod}_upl")
    if not files:
        return
    with st.expander(f"🖼️ Preview ({len(files)} files)", expanded=len(files) <= 5):
        cols = st.columns(5)
        for i, f in enumerate(files):
            with cols[i % 5]:
                try:
                    img = thumb(f)
                    if img is not None:
                        st.image(img, caption=f.name, width="stretch")
                    else:
                        st.caption(f"📄 {f.name}")
                except Exception as e:
                    st.error(f"Can't open {f.name}: {e}")

    processed = st.session_state[f"{mod}_processed"]
    c1, c2 = st.columns([1, 3])
    reprocess = c2.checkbox("Re-extract files already processed", key=f"{mod}_re")
    targets = files if reprocess else [f for f in files if f.name not in processed]
    if not c1.button(f"🔍 Extract ({len(targets)} file{'s' if len(targets) != 1 else ''})",
                     type="primary", disabled=not targets, key=f"{mod}_btn", width="stretch"):
        return

    bar = st.progress(0.0, "Reading…")
    new_rows = []
    for i, f in enumerate(targets, 1):
        try:
            rows, previews = read_file(f)
            if not rows:
                st.warning(f"No text found in **{f.name}** – retake the photo (flat, well lit, in focus).")
                continue
            recs, text = parse_fn(rows, f.name)
            new_rows.extend(recs)
            st.session_state[f"{mod}_images"][f.name] = previews
            st.session_state[f"{mod}_text"][f.name] = text
            processed.add(f.name)
        except Exception as e:
            st.error(f"❌ {f.name}: {e}")
        bar.progress(i / len(targets), f"Read {i}/{len(targets)}: {f.name}")

    if new_rows:
        current = st.session_state[f"{mod}_view"]  # keep manual edits
        if reprocess:
            current = current[~current["Source File"].isin([f.name for f in targets])]
        st.session_state[f"{mod}_base"] = pd.concat([current, pd.DataFrame(new_rows)], ignore_index=True)
        st.session_state[f"{mod}_ver"] += 1
        st.rerun()


def review(mod, column_config):
    edited = st.data_editor(
        st.session_state[f"{mod}_base"], num_rows="dynamic", width="stretch",
        height=430, column_config=column_config, key=f"{mod}_editor_{st.session_state[f'{mod}_ver']}",
    )
    st.session_state[f"{mod}_view"] = edited
    imgs = st.session_state[f"{mod}_images"]
    if imgs:
        with st.expander("🔎 Compare with the original photo / text"):
            name = st.selectbox("File", list(imgs.keys()), key=f"{mod}_cmp")
            a, b = st.columns([3, 2])
            for p in imgs[name]:
                a.image(p, width="stretch")
            b.text_area("Text read", st.session_state[f"{mod}_text"].get(name, ""), height=500,
                        key=f"{mod}_txt_{name}")
    return edited


# ============================================================================
# 8. TABS
# ============================================================================
tab_inv, tab_tm, tab_sum = st.tabs(["📦 Material Inventory", "🕒 Time & Material", "📊 Summary"])

with tab_inv:
    c1, c2 = st.columns(2)
    received = c1.date_input("Date Received (default)", date.today(), key="inv_recv")
    ordered_by = c2.text_input("Ordered By (if not on slip)", "", key="inv_ordby")
    st.subheader("1️⃣ Upload packing slips / delivery tickets / invoices")
    st.caption("One slip per photo/PDF works best. Take photos flat, well lit, whole page in frame.")
    upload_and_extract("inv", lambda rows, fn: parse_inventory(rows, fn, received, ordered_by, dayfirst),
                       "packing slips")
    st.subheader("2️⃣ Review & edit")
    st.caption("Check Qty, Product # and Vendor against the photo. Extension = Qty × Unit Price.")
    inv_df = clean_inv(review("inv", {
        "Doc Type": st.column_config.SelectboxColumn(options=DOC_TYPES),
        "Category": st.column_config.SelectboxColumn(options=CATEGORIES),
        "Qty": st.column_config.NumberColumn(format="%.2f"),
        "Unit Price": st.column_config.NumberColumn(format="$%.2f"),
        "Extension": st.column_config.NumberColumn(format="$%.2f"),
        "Description": st.column_config.TextColumn(width="large"),
        "Source File": st.column_config.TextColumn(disabled=True),
    }))
    st.subheader("3️⃣ Download")
    st.download_button("⬇️ Download Material Inventory Excel",
                       make_workbook(project, inv=inv_df, texts=st.session_state["inv_text"]),
                       f"Material_Inventory_{date.today():%Y%m%d}.xlsx",
                       type="primary", disabled=inv_df.empty, key="inv_dl")

with tab_tm:
    st.subheader("1️⃣ Upload T&M tickets / timesheets")
    st.caption("One ticket per photo/PDF. Printed names beside signatures are read; "
               "signatures themselves can't be verified – check the photo.")
    upload_and_extract("tm", lambda rows, fn: parse_tm(rows, fn, dayfirst), "T&M tickets")
    st.subheader("2️⃣ Review & edit")
    st.caption("Check worker names, hours and signatures. Total Hrs = Reg + OT. "
               "'verify' signatures are flagged red in Excel – type the name, or 'NOT SIGNED'.")
    tm_df = clean_tm(review("tm", {
        "Line Type": st.column_config.SelectboxColumn(options=LINE_TYPES),
        "Scope of Work": st.column_config.TextColumn(width="large"),
        "Reg Hrs": st.column_config.NumberColumn(format="%.2f"),
        "OT Hrs": st.column_config.NumberColumn(format="%.2f"),
        "Total Hrs": st.column_config.NumberColumn(format="%.2f"),
        "Qty": st.column_config.NumberColumn(format="%.2f"),
        "Unit Rate": st.column_config.NumberColumn(format="$%.2f"),
        "Amount": st.column_config.NumberColumn(format="$%.2f"),
        "Source File": st.column_config.TextColumn(disabled=True),
    }))
    st.subheader("3️⃣ Download")
    st.download_button("⬇️ Download T&M Tracking Excel",
                       make_workbook(project, tm=tm_df, texts=st.session_state["tm_text"]),
                       f"TM_Tracking_{date.today():%Y%m%d}.xlsx",
                       type="primary", disabled=tm_df.empty, key="tm_dl")

with tab_sum:
    st.subheader("📦 Material Inventory Summary")
    m = st.columns(4)
    m[0].metric("Line items", len(inv_df))
    m[1].metric("Documents", inv_df["Source File"].replace("", pd.NA).nunique())
    m[2].metric("Vendors", inv_df["Vendor"].replace("", pd.NA).nunique())
    m[3].metric("Total value", f"${pd.to_numeric(inv_df['Extension'], errors='coerce').fillna(0).sum():,.2f}")
    if not inv_df.empty:
        a, b = st.columns(2)
        a.markdown("**By Vendor**")
        a.dataframe(inv_df.groupby("Vendor").agg(Documents=("Source File", "nunique"),
                                                 Line_Items=("Description", "count"),
                                                 Total_Value=("Extension", "sum")).reset_index(),
                    hide_index=True, width="stretch",
                    column_config={"Total_Value": st.column_config.NumberColumn(format="$%.2f")})
        b.markdown("**By Category**")
        b.dataframe(inv_df.groupby("Category").agg(Line_Items=("Description", "count"),
                                                   Total_Qty=("Qty", "sum"),
                                                   Total_Value=("Extension", "sum")).reset_index(),
                    hide_index=True, width="stretch",
                    column_config={"Total_Value": st.column_config.NumberColumn(format="$%.2f")})

    st.divider()
    st.subheader("🕒 Time & Material Summary")
    tickets = tm_df.drop_duplicates(["Source File"]) if not tm_df.empty else tm_df
    lab = tm_df[tm_df["Line Type"] == "Labour"]
    unsigned = tickets[needs_sig(tickets["Trade Signed By"]) | needs_sig(tickets["Approved By (Bird)"])] \
        if not tickets.empty else tickets
    m = st.columns(5)
    m[0].metric("T&M tickets", len(tickets))
    m[1].metric("Trade companies", tm_df["Trade Company"].replace("", pd.NA).nunique())
    m[2].metric("Labour hours", f"{pd.to_numeric(lab['Total Hrs'], errors='coerce').fillna(0).sum():,.1f}")
    m[3].metric("Total amount", f"${pd.to_numeric(tm_df['Amount'], errors='coerce').fillna(0).sum():,.2f}")
    m[4].metric("Check signature", len(unsigned))
    if not tm_df.empty:
        rows = []
        for c, g in tm_df.groupby("Trade Company"):
            L = g["Line Type"]
            rows.append({"Trade Company": c, "Tickets": g["Source File"].nunique(),
                         "Workers": g.loc[L == "Labour", "Worker Name"].nunique(),
                         "Labour Hrs": g.loc[L == "Labour", "Total Hrs"].sum(),
                         "Labour $": g.loc[L == "Labour", "Amount"].sum(),
                         "Material $": g.loc[L == "Material", "Amount"].sum(),
                         "Equipment $": g.loc[L == "Equipment", "Amount"].sum(),
                         "Total $": g["Amount"].sum()})
        money = {c: st.column_config.NumberColumn(format="$%.2f")
                 for c in ("Labour $", "Material $", "Equipment $", "Total $")}
        st.markdown("**By Trade Company**")
        st.dataframe(pd.DataFrame(rows), hide_index=True, width="stretch", column_config=money)
        st.markdown("**By Ticket**")
        st.dataframe(tickets[["Work Date", "Trade Company", "T&M Ticket #", "Scope of Work",
                              "Trade Signed By", "Approved By (Bird)", "Signed Date"]],
                     hide_index=True, width="stretch")
        if not unsigned.empty:
            st.warning(f"⚠ {len(unsigned)} ticket(s) need a signature check – confirm on the photo "
                       "before submitting for change order pricing.")

    st.divider()
    st.download_button("⬇️ Download Combined Excel (Inventory + T&M + Summaries)",
                       make_workbook(project, inv=inv_df, tm=tm_df,
                                     texts={**st.session_state["inv_text"], **st.session_state["tm_text"]}),
                       f"MMCC_Inventory_TM_{date.today():%Y%m%d}.xlsx",
                       type="primary", disabled=inv_df.empty and tm_df.empty, key="all_dl")
